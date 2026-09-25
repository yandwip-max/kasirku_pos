"""Backup data (JSON snapshot + real Excel exports): store-scoped, owner-only.

Regression coverage for BUG 2 ("Gagal mengunduh stok imei" / "hasilnya kacau"):
the old /backup/csv/{dataset} route (missing decorator -> 404) was replaced by
GET /backup/xlsx/{products|units|transactions} which returns a real .xlsx
workbook (pandas + openpyxl) with Indonesian headers and one value per column.
"""
import io
import json
import uuid

import httpx
import openpyxl

from tests.helpers import KASIR, PEMILIK, auth_headers


def _create_fixture_product(client: httpx.Client, headers: dict) -> str:
    suffix = uuid.uuid4().hex[:8]
    product = client.post(
        "/products",
        json={
            "name": f"tscheck-backup-prod-{suffix}",
            "type": "aksesoris",
            "category": "Aksesoris",
            "sell_price": 12000,
            "cost_price": 8000,
            "stock_qty": 5,
            "min_stock": 1,
        },
        headers=headers,
    )
    assert product.status_code == 201, product.text
    return product.json()["id"]


def _create_fixture_unit(client: httpx.Client, headers: dict) -> tuple[str, str]:
    """A handphone product + one IMEI unit, so the units sheet has a real row."""
    suffix = uuid.uuid4().hex[:8]
    product = client.post(
        "/products",
        json={
            "name": f"tscheck-backup-hp-{suffix}",
            "type": "handphone",
            "category": "Handphone",
            "sell_price": 1500000,
            "cost_price": 1200000,
            "stock_qty": 0,
            "min_stock": 0,
        },
        headers=headers,
    )
    assert product.status_code == 201, product.text
    product_id = product.json()["id"]

    imei = f"86{uuid.uuid4().int % 10**13:013d}"
    unit = client.post(
        f"/products/{product_id}/units",
        json={"imei": imei, "color": "Hitam", "capacity": "128GB", "cost_price": 1200000, "sell_price": 1500000},
        headers=headers,
    )
    assert unit.status_code == 201, unit.text
    return product_id, imei


def test_backup_json_export_contains_counts_and_own_store_data_only(client: httpx.Client):
    headers = auth_headers(client, PEMILIK)
    product_id = _create_fixture_product(client, headers)

    resp = client.get("/backup/export", headers=headers)
    assert resp.status_code == 200, resp.text
    assert "attachment" in resp.headers.get("content-disposition", "")
    assert resp.headers.get("content-disposition", "").endswith('.json"')

    snapshot = json.loads(resp.text)
    assert snapshot["app"] == "KasirKu"
    assert "counts" in snapshot
    for key in ("products", "product_units", "transactions", "activity_logs"):
        assert key in snapshot["data"]
        assert snapshot["counts"][key] == len(snapshot["data"][key])

    product_ids = [p["id"] for p in snapshot["data"]["products"]]
    assert product_id in product_ids
    store_ids = {p.get("store_id") for p in snapshot["data"]["products"] if "store_id" in p}
    assert store_ids in (set(), {snapshot["store_id"]})


def test_backup_xlsx_units_has_indonesian_headers_separate_columns_and_numeric_price(client: httpx.Client):
    headers = auth_headers(client, PEMILIK)
    product_id, imei = _create_fixture_unit(client, headers)

    resp = client.get("/backup/xlsx/units", headers=headers)
    assert resp.status_code == 200, resp.text
    assert resp.headers.get("content-type") == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    disposition = resp.headers.get("content-disposition", "")
    assert "attachment" in disposition
    import re

    match = re.search(r'filename="(kasirku-units-\d{4}-\d{2}-\d{2}\.xlsx)"', disposition)
    assert match, disposition

    workbook = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert workbook.sheetnames == ["Stok IMEI"]
    sheet = workbook["Stok IMEI"]
    header_row = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert header_row == [
        "Nama Produk", "IMEI", "Warna", "Kapasitas", "Status",
        "Harga Modal", "Harga Jual", "Tanggal Masuk", "Tanggal Terjual",
    ]

    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    matching = [r for r in rows if r[1] == imei]
    assert len(matching) == 1, f"fixture IMEI {imei} not found as its own cell in {len(rows)} rows"
    row = matching[0]
    assert row[4] == "Tersedia"  # status label, in its own column (not concatenated)
    assert isinstance(row[6], (int, float))  # Harga Jual is a real number, not text
    assert row[6] == 1500000


def test_backup_xlsx_products_and_transactions_have_indonesian_headers(client: httpx.Client):
    headers = auth_headers(client, PEMILIK)
    product_id = _create_fixture_product(client, headers)

    products_resp = client.get("/backup/xlsx/products", headers=headers)
    assert products_resp.status_code == 200, products_resp.text
    products_wb = openpyxl.load_workbook(io.BytesIO(products_resp.content))
    assert products_wb.sheetnames == ["Produk"]
    products_header = [c.value for c in next(products_wb["Produk"].iter_rows(min_row=1, max_row=1))]
    assert products_header == [
        "Nama Produk", "Tipe", "Folder / Kategori", "Merek", "SKU",
        "Harga Modal", "Harga Jual", "Harga Grosir", "Stok", "Stok Minimum",
    ]
    product_names = [r[0] for r in products_wb["Produk"].iter_rows(min_row=2, values_only=True)]
    assert any(name and name.startswith("tscheck-backup-prod-") for name in product_names)

    tx_resp = client.get("/backup/xlsx/transactions", headers=headers)
    assert tx_resp.status_code == 200, tx_resp.text
    tx_wb = openpyxl.load_workbook(io.BytesIO(tx_resp.content))
    assert tx_wb.sheetnames == ["Transaksi"]
    tx_header = [c.value for c in next(tx_wb["Transaksi"].iter_rows(min_row=1, max_row=1))]
    assert tx_header == [
        "No. Struk", "Tanggal", "Status", "Kasir", "Pembeli", "Pembayaran",
        "Subtotal", "Diskon", "Total", "Laba", "Rincian Barang",
    ]


def test_backup_unknown_dataset_404_kasir_forbidden_403_no_token_401_old_csv_route_404(client: httpx.Client):
    headers = auth_headers(client, PEMILIK)

    unknown = client.get("/backup/xlsx/not-a-real-dataset", headers=headers)
    assert unknown.status_code == 404, unknown.text

    kasir_headers = auth_headers(client, KASIR)
    forbidden_json = client.get("/backup/export", headers=kasir_headers)
    assert forbidden_json.status_code == 403, forbidden_json.text

    forbidden_xlsx = client.get("/backup/xlsx/units", headers=kasir_headers)
    assert forbidden_xlsx.status_code == 403, forbidden_xlsx.text

    no_token = client.get("/backup/xlsx/units")
    assert no_token.status_code == 401, no_token.text

    # BUG 2 regression: the old CSV route must be gone, not silently 200
    old_route = client.get("/backup/csv/units", headers=headers)
    assert old_route.status_code == 404, old_route.text
